import openpyxl
import io


def parse_excel(file_bytes: bytes) -> str:
    """
    Extracts all text from an Excel file (.xlsx).
    Reads every sheet and every cell, returns a plain text representation.
    """
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    text_parts = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        text_parts.append(f"Sheet: {sheet_name}")

        for row in sheet.iter_rows(values_only=True):
            row_values = [str(cell) for cell in row if cell is not None]
            if row_values:
                text_parts.append(" | ".join(row_values))

    return "\n".join(text_parts)